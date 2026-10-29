"""
分析 Datasets Fusion Excel 文件的数据结构 - 输出到文件
"""
import openpyxl
from collections import defaultdict
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook(r"C:\Users\Cupshe\Downloads\Datasets Fusion 2026-09-17 16-30-37.xlsx", read_only=True)

output = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    
    header = rows[0]
    output.append(f"Sheet: {sheet_name}")
    output.append(f"列数: {len(header)}")
    output.append(f"数据行数: {len(rows) - 1}")
    output.append(f"字段名: {list(header)}")
    output.append("")
    
    # 打印前5行数据
    output.append("前5行数据:")
    for i, row in enumerate(rows[1:6], 1):
        output.append(f"  Row {i}:")
        for j, val in enumerate(row):
            output.append(f"    {header[j]}: {val}")
        output.append("")
    
    # 各列统计
    output.append("各列统计:")
    for col_idx in range(len(header)):
        values = [row[col_idx] for row in rows[1:] if row[col_idx] is not None]
        null_count = sum(1 for row in rows[1:] if row[col_idx] is None)
        unique_count = len(set(str(v) for v in values))
        if values:
            output.append(f"  [{header[col_idx]}] 非空:{len(values)}, 空:{null_count}, 唯一值:{unique_count}, 类型:{type(values[0]).__name__}")
            if unique_count <= 15:
                output.append(f"    值分布: {dict((str(k), sum(1 for v in values if str(v)==str(k))) for k in set(str(v) for v in values))}")
            else:
                output.append(f"    样例: {values[:3]}")
        else:
            output.append(f"  [{header[col_idx]}] 全部为空")

wb.close()

with open('dataset_analysis.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))
print("Done! Output saved to dataset_analysis.txt")
